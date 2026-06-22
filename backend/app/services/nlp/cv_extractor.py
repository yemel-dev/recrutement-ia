"""
cv_extractor.py
---------------
Membre 2 - NLP/IA | Recrutement Intelligent | Université de Dschang 2024-2025

Rôle : Extraire le texte brut d'un CV au format PDF ou DOCX.
C'est la première étape du pipeline NLP — sans texte, rien ne peut être analysé.

Dépendances :
    - pymupdf (import fitz)  : lecture des fichiers PDF
    - python-docx            : lecture des fichiers DOCX
"""

import fitz                         # PyMuPDF — lecture PDF
from docx import Document           # python-docx — lecture DOCX
from pathlib import Path


# ──────────────────────────────────────────────────────────────────────────────
# Fonctions privées (usage interne uniquement)
# ──────────────────────────────────────────────────────────────────────────────

def _extraire_texte_pdf(chemin_fichier: str) -> str:
    """
    Ouvre un fichier PDF et extrait tout le texte page par page.

    Args:
        chemin_fichier : chemin absolu ou relatif vers le fichier .pdf

    Returns:
        Texte brut complet du PDF (toutes pages concaténées).

    Raises:
        FileNotFoundError : si le fichier n'existe pas.
        ValueError        : si le PDF est vide ou illisible.
    """
    chemin = Path(chemin_fichier)

    if not chemin.exists():
        raise FileNotFoundError(f"Fichier PDF introuvable : {chemin_fichier}")

    texte_pages = []

    # fitz.open() ouvre le PDF — chaque élément de la boucle est une page
    with fitz.open(str(chemin)) as document_pdf:

        if document_pdf.page_count == 0:
            raise ValueError(f"Le PDF est vide (0 pages) : {chemin_fichier}")

        for numero_page, page in enumerate(document_pdf, start=1):
            # get_text("text") extrait le texte brut de la page
            texte_page = page.get_text("text")
            texte_pages.append(texte_page)

    texte_complet = "\n".join(texte_pages).strip()

    if not texte_complet:
        raise ValueError(
            f"Aucun texte extrait du PDF : {chemin_fichier}. "
            "Le fichier est peut-être un scan (image). "
            "L'OCR n'est pas encore supporté."
        )

    return texte_complet


def _extraire_texte_docx(chemin_fichier: str) -> str:
    """
    Ouvre un fichier DOCX et extrait tout le texte paragraphe par paragraphe.

    Args:
        chemin_fichier : chemin absolu ou relatif vers le fichier .docx

    Returns:
        Texte brut complet du DOCX.

    Raises:
        FileNotFoundError : si le fichier n'existe pas.
        ValueError        : si le DOCX est vide.
    """
    chemin = Path(chemin_fichier)

    if not chemin.exists():
        raise FileNotFoundError(f"Fichier DOCX introuvable : {chemin_fichier}")

    # Document() charge le fichier Word
    document_word = Document(str(chemin))

    # Chaque paragraphe est un bloc de texte (titre, corps, liste, etc.)
    paragraphes = [
        paragraphe.text
        for paragraphe in document_word.paragraphs
        if paragraphe.text.strip()   # on ignore les paragraphes vides
    ]

    texte_complet = "\n".join(paragraphes).strip()

    if not texte_complet:
        raise ValueError(f"Aucun texte extrait du DOCX : {chemin_fichier}")

    return texte_complet


# ──────────────────────────────────────────────────────────────────────────────
# Fonction publique principale (utilisée par le pipeline NLP)
# ──────────────────────────────────────────────────────────────────────────────

def extraire_texte_cv(chemin_fichier: str) -> str:
    """
    Fonction principale du module.
    Détecte automatiquement le format du fichier et extrait son texte.

    Cette fonction est appelée par nlp_pipeline.py — c'est le point d'entrée
    unique pour toute extraction de texte de CV.

    Args:
        chemin_fichier : chemin vers le CV (PDF ou DOCX).

    Returns:
        Texte brut extrait du CV, prêt à être envoyé au prétraitement NLP.

    Raises:
        FileNotFoundError : si le fichier n'existe pas.
        ValueError        : si le format n'est pas supporté ou le fichier vide.

    Exemple:
        >>> texte = extraire_texte_cv("uploads/cv_jean_dupont.pdf")
        >>> print(texte[:100])
        'Jean Dupont\\nIngénieur en Intelligence Artificielle\\n...'
    """
    chemin = Path(chemin_fichier)
    extension = chemin.suffix.lower()   # ".pdf", ".docx", etc.

    if extension == ".pdf":
        return _extraire_texte_pdf(chemin_fichier)

    elif extension == ".docx":
        return _extraire_texte_docx(chemin_fichier)

    else:
        raise ValueError(
            f"Format de fichier non supporté : '{extension}'. "
            "Seuls les formats PDF (.pdf) et Word (.docx) sont acceptés."
        )
